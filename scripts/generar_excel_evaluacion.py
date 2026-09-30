"""
Generador de la plantilla de evaluación del reto de pedidos — CGAO Skills.

Produce un libro de Excel (.xlsx) con:
  • Hoja «Evaluación»: 12 criterios del reto con M (Medición), J (Juicio)
    y P (Proceso), para un total de 100 puntos.
  • Entradas vacías para tres aprendices, totales ponderados y firmas.

Requiere: openpyxl >= 3.1
"""

from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Paleta de colores
# ---------------------------------------------------------------------------
SENA_GREEN = "005A2B"
DARK_GREEN = "003D1E"
GOLD = "D4A017"
LIGHT_GOLD = "FFF8DC"
LIGHT_GREEN = "E8F5E9"
WHITE = "FFFFFF"
LIGHT_GRAY = "F5F5F5"
MEDIUM_GRAY = "BDBDBD"
DARK_GRAY = "424242"
BLUE_MED = "1565C0"
BLUE_LIGHT = "E3F2FD"
ORANGE_MED = "E65100"
ORANGE_LIGHT = "FFF3E0"

# ---------------------------------------------------------------------------
# Estilos reutilizables
# ---------------------------------------------------------------------------
THIN_BORDER = Border(
    left=Side(style="thin", color=MEDIUM_GRAY),
    right=Side(style="thin", color=MEDIUM_GRAY),
    top=Side(style="thin", color=MEDIUM_GRAY),
    bottom=Side(style="thin", color=MEDIUM_GRAY),
)

HEADER_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=SENA_GREEN, size=16)
SUBTITLE_FONT = Font(name="Calibri", bold=True, color=DARK_GRAY, size=12)
SECTION_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
NORMAL_FONT = Font(name="Calibri", size=10, color=DARK_GRAY)
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color=DARK_GRAY)
SIGNATURE_FONT = Font(name="Calibri", size=10, color=DARK_GRAY, italic=True)
SMALL_FONT = Font(name="Calibri", size=9, color="757575")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)

FILL_HEADER_GREEN = PatternFill("solid", fgColor=SENA_GREEN)
FILL_HEADER_BLUE = PatternFill("solid", fgColor=BLUE_MED)
FILL_HEADER_ORANGE = PatternFill("solid", fgColor=ORANGE_MED)
FILL_GOLD = PatternFill("solid", fgColor=LIGHT_GOLD)
FILL_LIGHT_GREEN = PatternFill("solid", fgColor=LIGHT_GREEN)
FILL_LIGHT_BLUE = PatternFill("solid", fgColor=BLUE_LIGHT)
FILL_LIGHT_ORANGE = PatternFill("solid", fgColor=ORANGE_LIGHT)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)
FILL_LIGHT_GRAY = PatternFill("solid", fgColor=LIGHT_GRAY)

# ---------------------------------------------------------------------------
# Criterios de la rúbrica maestra del reto de pedidos
# ---------------------------------------------------------------------------
CRITERIOS_MEDICION = [
    {
        "id": "M1",
        "modulo": "A – Fundamentos",
        "aspecto": "Reglas y cálculo del pedido",
        "descripcion": "Alias de 2 a 40 caracteres; producto del catálogo; cantidad entera de 1 a 100; subtotal igual a cantidad por precio unitario",
        "pts_max": 7.0,
    },
    {
        "id": "M2",
        "modulo": "B – Construcción",
        "aspecto": "Registro del pedido",
        "descripcion": "Código consecutivo y campos correctos; 2 bocadillos suman $15.000 y 3 panelitas suman $15.000",
        "pts_max": 15.0,
    },
    {
        "id": "M3",
        "modulo": "B – Construcción",
        "aspecto": "Validación de datos",
        "descripcion": "Rechaza alias vacío o fuera del rango y cantidad no entera o fuera del rango; solo permite productos del catálogo",
        "pts_max": 12.0,
    },
    {
        "id": "M4",
        "modulo": "B – Construcción",
        "aspecto": "Consulta de pedidos",
        "descripcion": "Lista todos los campos; busca alias sin distinguir mayúsculas; filtra por estado",
        "pts_max": 10.0,
    },
    {
        "id": "M5",
        "modulo": "B – Construcción",
        "aspecto": "Estado, resumen y persistencia",
        "descripcion": "Cambia el estado; muestra la cantidad y el valor de pedidos entregados; conserva pedidos y estados en el almacenamiento local al recargar",
        "pts_max": 13.0,
    },
    {
        "id": "M6",
        "modulo": "C – Ordenamiento",
        "aspecto": "Ordenamiento por subtotal",
        "descripcion": "Orden descendente de mayor a menor subtotal; desempate por código menor; ajuste completado en 30 minutos",
        "pts_max": 10.0,
    },
    {
        "id": "M7",
        "modulo": "D – Sostenibilidad y cierre",
        "aspecto": "Cierre y uso eficiente de la estación",
        "descripcion": "Al finalizar, cierra los programas que ya no necesita, retira los archivos temporales de prueba creados para el reto y apaga el monitor o el equipo si no se va a usar de inmediato",
        "pts_max": 2.0,
    },
]

CRITERIOS_JUZGAMIENTO = [
    {
        "id": "J1",
        "modulo": "A – Fundamentos",
        "aspecto": "Planeación técnica",
        "escala": "0 = Sin evidencia · 1 = Flujo parcial · 2 = Incluye acciones requeridas · 3 = Completo y fácil de seguir",
        "pts_max": 8.0,
    },
    {
        "id": "J2",
        "modulo": "B – Construcción",
        "aspecto": "Interfaz de la aplicación",
        "escala": "0 = Inoperable · 1 = Fallas de uso · 2 = Tareas claras · 3 = Uso fácil y errores claros",
        "pts_max": 10.0,
    },
    {
        "id": "J3",
        "modulo": "C – Ordenamiento",
        "aspecto": "Explicación de la comprobación",
        "escala": "0 = Sin explicación o incorrecta · 1 = Menciona resultado · 2 = Describe pruebas · 3 = Explica pruebas y resultado",
        "pts_max": 5.0,
    },
]

CRITERIOS_PROCESO = [
    {
        "id": "P1",
        "modulo": "D – Sostenibilidad y cierre",
        "aspecto": "Ergonomía y orden del puesto",
        "descripcion": "Durante el reto mantiene una postura adecuada, ordena los elementos de trabajo y deja despejada la estación",
        "pts_max": 4.0,
    },
    {
        "id": "P2",
        "modulo": "D – Sostenibilidad y cierre",
        "aspecto": "Separación de residuos",
        "descripcion": "Deposita cada residuo en el recipiente identificado con su color: blanco para residuos aprovechables limpios y secos, verde para residuos orgánicos aprovechables y negro para residuos no aprovechables. Si no genera residuos, deja la estación limpia",
        "pts_max": 4.0,
    },
]

PARTICIPANTES = [
    {"nombre": "", "programa": "", "ficha": ""},
    {"nombre": "", "programa": "", "ficha": ""},
    {"nombre": "", "programa": "", "ficha": ""},
]

JUECES = [
    {"nombre": "Instructor evaluador 1", "cargo": "Firma y observaciones"},
    {"nombre": "Instructor evaluador 2", "cargo": "Firma y observaciones"},
    {"nombre": "Instructor evaluador 3", "cargo": "Firma y observaciones"},
]


def _apply_style(cell, font=None, fill=None, alignment=None, border=None,
                 number_format=None):
    """Aplica estilos a una celda de forma compacta."""
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format


def _write_row(ws, row, col_start, values, font=None, fill=None,
               alignment=None, border=THIN_BORDER, number_format=None):
    """Escribe una fila de valores y devuelve la fila siguiente."""
    for i, val in enumerate(values):
        cell = ws.cell(row=row, column=col_start + i, value=val)
        _apply_style(cell, font=font, fill=fill, alignment=alignment,
                     border=border, number_format=number_format)
    return row + 1


def _merge_and_style(ws, start_row, start_col, end_row, end_col, value,
                     font, fill, alignment=CENTER, border=THIN_BORDER):
    """Combina celdas y aplica estilo."""
    ws.merge_cells(
        start_row=start_row,
        start_column=start_col,
        end_row=end_row,
        end_column=end_col,
    )
    cell = ws.cell(row=start_row, column=start_col, value=value)
    _apply_style(cell, font=font, fill=fill, alignment=alignment, border=border)
    # Aplicar borde a todas las celdas del rango combinado
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(row=r, column=c).border = border


def generar_excel(output_path: Path) -> Path:
    """Genera el libro de Excel y lo guarda en *output_path*."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evaluación"

    # Configuración de página
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(
        fitToPage=True,
    )
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = "F1"
    ws.sheet_properties.tabColor = SENA_GREEN

    # Anchos de columna
    col_widths = {
        1: 6,    # ID
        2: 22,   # Módulo
        3: 40,   # Aspecto / Criterio
        4: 44,   # Descripción / Escala
        5: 10,   # Pts Máx
        6: 14,   # Participante 1
        7: 14,   # Participante 2
        8: 14,   # Participante 3
    }
    for col_num, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width

    num_participantes = len(PARTICIPANTES)
    total_cols = 5 + num_participantes  # 5 fijas + N participantes

    row = 1

    # ===================================================================
    # ENCABEZADO INSTITUCIONAL
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "CENTRO DE GESTIÓN AGROEMPRESARIAL DEL ORIENTE — CGAO",
        Font(name="Calibri", bold=True, color=SENA_GREEN, size=14),
        FILL_WHITE,
        CENTER,
    )
    row += 1

    _merge_and_style(
        ws, row, 1, row, total_cols,
        "PLANTILLA DE EVALUACIÓN DEL RETO DE REGISTRO DE PEDIDOS",
        Font(name="Calibri", bold=True, color=DARK_GRAY, size=13),
        FILL_LIGHT_GREEN,
        CENTER,
    )
    row += 1

    _merge_and_style(
        ws, row, 1, row, total_cols,
        "Programa: Análisis y Desarrollo de Software  |  Subsede Vélez  |  Ejemplo formativo: confirme condiciones con el instructor antes de usarlo.",
        SMALL_FONT,
        FILL_WHITE,
        CENTER,
    )
    row += 2  # Línea en blanco

    # ===================================================================
    # DATOS DE PARTICIPANTES
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "DATOS DE PARTICIPANTES",
        SECTION_FONT,
        FILL_HEADER_GREEN,
        CENTER,
    )
    row += 1

    headers_part = ["#", "Nombre completo", "Programa de formación",
                    "N.° de ficha", "", "", "", ""]
    part_header_values = headers_part[:total_cols]
    _write_row(ws, row, 1, part_header_values[:4],
               font=BOLD_FONT, fill=FILL_LIGHT_GREEN, alignment=CENTER)
    row += 1

    for idx, p in enumerate(PARTICIPANTES, 1):
        vals = [idx, p["nombre"], p["programa"], p["ficha"]]
        _write_row(ws, row, 1, vals,
                   font=NORMAL_FONT, fill=FILL_WHITE, alignment=CENTER)
        for col in range(2, 5):
            ws.cell(row=row, column=col).fill = FILL_GOLD
        row += 1

    row += 1  # Separador

    # ===================================================================
    # INSTRUCCIONES DE CALIFICACIÓN
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        ("Complete los campos amarillos. En M y P registre 0 o el máximo del criterio. "
         "En J registre un nivel entero de 0 a 3; "
         "la planilla lo convierte en puntos con nivel ÷ 3 × puntaje máximo."),
        Font(name="Calibri", size=9, bold=True, color=DARK_GRAY),
        FILL_GOLD,
        CENTER,
    )
    row += 1

    _merge_and_style(
        ws, row, 1, row, total_cols,
        "Máximos: M = 69 puntos · J = 23 puntos · P = 8 puntos · Total = 100 puntos",
        Font(name="Calibri", size=9, italic=True, color=DARK_GRAY),
        FILL_WHITE,
        CENTER,
    )
    row += 2

    # ===================================================================
    # SECCIÓN 1: CRITERIOS DE MEDICIÓN (M)
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "SECCIÓN 1 — CRITERIOS DE MEDICIÓN (M)  ·  Mínimo 60 % del puntaje total",
        SECTION_FONT,
        FILL_HEADER_BLUE,
        CENTER,
    )
    row += 1

    # Encabezado de tabla
    med_headers = ["ID", "Módulo", "Aspecto evaluado",
                   "Descripción de cumplimiento", "Pts máx"]
    for p_idx, _p in enumerate(PARTICIPANTES, 1):
        med_headers.append(f"Puntaje {p_idx}")
    _write_row(ws, row, 1, med_headers,
               font=Font(name="Calibri", bold=True, color=WHITE, size=10),
               fill=FILL_HEADER_BLUE,
               alignment=CENTER)
    header_row_m = row
    row += 1

    first_data_row_m = row
    for crit in CRITERIOS_MEDICION:
        vals = [crit["id"], crit["modulo"], crit["aspecto"],
                crit["descripcion"], crit["pts_max"]]
        vals.extend([None] * num_participantes)
        bg = FILL_WHITE if (row - first_data_row_m) % 2 == 0 else FILL_LIGHT_BLUE
        _write_row(ws, row, 1, vals,
                   font=NORMAL_FONT, fill=bg, alignment=CENTER)
        # La columna de descripción alineada a la izquierda
        ws.cell(row=row, column=3).alignment = LEFT_WRAP
        ws.cell(row=row, column=4).alignment = LEFT_WRAP
        if crit["id"] == "M7":
            ws.row_dimensions[row].height = 42
        row += 1
    last_data_row_m = row - 1

    medicion_validation = DataValidation(
        type="custom",
        formula1=f"OR(F{first_data_row_m}=0,F{first_data_row_m}=$E{first_data_row_m})",
        allow_blank=True,
    )
    medicion_validation.error = "Registre 0 o el puntaje máximo del criterio."
    medicion_validation.errorTitle = "Puntaje no válido"
    medicion_validation.prompt = "Use 0 o el máximo de puntos de la fila."
    medicion_validation.promptTitle = "Criterio de medición"
    ws.add_data_validation(medicion_validation)
    medicion_validation.add(f"F{first_data_row_m}:H{last_data_row_m}")
    for col in range(6, total_cols + 1):
        for current_row in range(first_data_row_m, last_data_row_m + 1):
            ws.cell(row=current_row, column=col).fill = FILL_GOLD
            ws.cell(row=current_row, column=col).font = BOLD_FONT

    # Subtotal Medición
    subtotal_vals = ["", "", "", "SUBTOTAL MEDICIÓN (M)", ""]
    pts_max_sum_m = sum(c["pts_max"] for c in CRITERIOS_MEDICION)
    subtotal_vals[4] = pts_max_sum_m
    for p_idx in range(num_participantes):
        col_letter = get_column_letter(6 + p_idx)
        formula = (
            f'=IF(COUNT({col_letter}{first_data_row_m}:{col_letter}{last_data_row_m})=0,"",'
            f"SUM({col_letter}{first_data_row_m}:{col_letter}{last_data_row_m}))"
        )
        subtotal_vals.append(formula)
    for i, val in enumerate(subtotal_vals):
        cell = ws.cell(row=row, column=1 + i, value=val)
        _apply_style(cell, font=BOLD_FONT, fill=FILL_LIGHT_BLUE,
                     alignment=CENTER, border=THIN_BORDER,
                     number_format="0.0")
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center", wrap_text=True,
    )
    subtotal_row_m = row
    row += 2

    # ===================================================================
    # SECCIÓN 2: CRITERIOS DE JUZGAMIENTO (J)
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "SECCIÓN 2 — CRITERIOS DE JUZGAMIENTO (J)  ·  Máximo 30 % del puntaje total",
        SECTION_FONT,
        FILL_HEADER_ORANGE,
        CENTER,
    )
    row += 1

    judg_headers = ["ID", "Módulo", "Aspecto evaluado",
                    "Escala de valoración (0-3)", "Pts máx"]
    for p_idx, _p in enumerate(PARTICIPANTES, 1):
        judg_headers.append(f"Nivel {p_idx} (0–3)")
    _write_row(ws, row, 1, judg_headers,
               font=Font(name="Calibri", bold=True, color=WHITE, size=10),
               fill=FILL_HEADER_ORANGE,
               alignment=CENTER)
    row += 1

    first_data_row_j = row
    for crit in CRITERIOS_JUZGAMIENTO:
        vals = [crit["id"], crit["modulo"], crit["aspecto"],
                crit["escala"], crit["pts_max"]]
        vals.extend([None] * num_participantes)
        bg = FILL_WHITE if (row - first_data_row_j) % 2 == 0 else FILL_LIGHT_ORANGE
        _write_row(ws, row, 1, vals,
                   font=NORMAL_FONT, fill=bg, alignment=CENTER)
        ws.cell(row=row, column=3).alignment = LEFT_WRAP
        ws.cell(row=row, column=4).alignment = LEFT_WRAP
        row += 1
    last_data_row_j = row - 1

    juzgamiento_validation = DataValidation(
        type="whole",
        operator="between",
        formula1=0,
        formula2=3,
        allow_blank=True,
    )
    juzgamiento_validation.error = "Registre un nivel entero entre 0 y 3."
    juzgamiento_validation.errorTitle = "Nivel no válido"
    juzgamiento_validation.prompt = "La planilla convierte el nivel a puntos ponderados."
    juzgamiento_validation.promptTitle = "Criterio de juicio"
    ws.add_data_validation(juzgamiento_validation)
    juzgamiento_validation.add(f"F{first_data_row_j}:H{last_data_row_j}")
    for col in range(6, total_cols + 1):
        for current_row in range(first_data_row_j, last_data_row_j + 1):
            ws.cell(row=current_row, column=col).fill = FILL_GOLD
            ws.cell(row=current_row, column=col).font = BOLD_FONT

    # Subtotal Juzgamiento
    subtotal_vals_j = ["", "", "", "SUBTOTAL JUZGAMIENTO (J)", ""]
    pts_max_sum_j = sum(c["pts_max"] for c in CRITERIOS_JUZGAMIENTO)
    subtotal_vals_j[4] = pts_max_sum_j
    for p_idx in range(num_participantes):
        col_letter = get_column_letter(6 + p_idx)
        formula = (
            f'=IF(COUNT({col_letter}{first_data_row_j}:{col_letter}{last_data_row_j})=0,"",'
            f"SUMPRODUCT({col_letter}{first_data_row_j}:{col_letter}{last_data_row_j},"
            f"$E${first_data_row_j}:$E${last_data_row_j})/3)"
        )
        subtotal_vals_j.append(formula)
    for i, val in enumerate(subtotal_vals_j):
        cell = ws.cell(row=row, column=1 + i, value=val)
        _apply_style(cell, font=BOLD_FONT, fill=FILL_LIGHT_ORANGE,
                     alignment=CENTER, border=THIN_BORDER,
                     number_format="0.0")
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center", wrap_text=True,
    )
    subtotal_row_j = row
    row += 2

    # ===================================================================
    # SECCIÓN 3: CRITERIOS DE PROCESO (P)
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "SECCIÓN 3 — CRITERIOS DE PROCESO (P)  ·  Máximo 10 % del puntaje total",
        SECTION_FONT,
        FILL_HEADER_GREEN,
        CENTER,
    )
    row += 1

    proc_headers = ["ID", "Módulo", "Aspecto evaluado",
                    "Descripción de cumplimiento", "Pts máx"]
    for p_idx, _p in enumerate(PARTICIPANTES, 1):
        proc_headers.append(f"Puntaje {p_idx}")
    _write_row(ws, row, 1, proc_headers,
               font=Font(name="Calibri", bold=True, color=WHITE, size=10),
               fill=FILL_HEADER_GREEN,
               alignment=CENTER)
    row += 1

    first_data_row_p = row
    for crit in CRITERIOS_PROCESO:
        vals = [crit["id"], crit["modulo"], crit["aspecto"],
                crit["descripcion"], crit["pts_max"]]
        vals.extend([None] * num_participantes)
        _write_row(ws, row, 1, vals,
                   font=NORMAL_FONT, fill=FILL_WHITE, alignment=CENTER)
        ws.cell(row=row, column=3).alignment = LEFT_WRAP
        ws.cell(row=row, column=4).alignment = LEFT_WRAP
        if crit["id"] == "P2":
            ws.row_dimensions[row].height = 68
        row += 1
    last_data_row_p = row - 1

    proceso_validation = DataValidation(
        type="custom",
        formula1=f"OR(F{first_data_row_p}=0,F{first_data_row_p}=$E{first_data_row_p})",
        allow_blank=True,
    )
    proceso_validation.error = "Registre 0 o el puntaje máximo del criterio."
    proceso_validation.errorTitle = "Puntaje no válido"
    proceso_validation.prompt = "Use 0 o el máximo de puntos de la fila."
    proceso_validation.promptTitle = "Criterio de proceso"
    ws.add_data_validation(proceso_validation)
    proceso_validation.add(f"F{first_data_row_p}:H{last_data_row_p}")
    for col in range(6, total_cols + 1):
        for current_row in range(first_data_row_p, last_data_row_p + 1):
            ws.cell(row=current_row, column=col).fill = FILL_GOLD
            ws.cell(row=current_row, column=col).font = BOLD_FONT

    # Subtotal Proceso
    subtotal_vals_p = ["", "", "", "SUBTOTAL PROCESO (P)", ""]
    pts_max_sum_p = sum(c["pts_max"] for c in CRITERIOS_PROCESO)
    subtotal_vals_p[4] = pts_max_sum_p
    for p_idx in range(num_participantes):
        col_letter = get_column_letter(6 + p_idx)
        formula = (
            f'=IF(COUNT({col_letter}{first_data_row_p}:{col_letter}{last_data_row_p})=0,"",'
            f"SUM({col_letter}{first_data_row_p}:{col_letter}{last_data_row_p}))"
        )
        subtotal_vals_p.append(formula)
    for i, val in enumerate(subtotal_vals_p):
        cell = ws.cell(row=row, column=1 + i, value=val)
        _apply_style(cell, font=BOLD_FONT, fill=FILL_LIGHT_GREEN,
                     alignment=CENTER, border=THIN_BORDER,
                     number_format="0.0")
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center", wrap_text=True,
    )
    subtotal_row_p = row
    row += 2

    # ===================================================================
    # RESUMEN GENERAL Y PUNTAJE FINAL
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "RESUMEN GENERAL — PUNTAJE FINAL",
        Font(name="Calibri", bold=True, color=WHITE, size=12),
        PatternFill("solid", fgColor=DARK_GREEN),
        CENTER,
    )
    row += 1

    # Encabezados del resumen
    resumen_headers = ["", "", "", "Concepto", "Pts máx"]
    for p_idx, _p in enumerate(PARTICIPANTES, 1):
        resumen_headers.append(f"Aprendiz {p_idx}")
    _write_row(ws, row, 1, resumen_headers,
               font=Font(name="Calibri", bold=True, color=WHITE, size=10),
               fill=PatternFill("solid", fgColor=DARK_GREEN),
               alignment=CENTER)
    row += 1

    pts_total_max = pts_max_sum_m + pts_max_sum_j + pts_max_sum_p

    # Fila: Total Medición
    r_vals = ["", "", "", "Total Medición (M)", pts_max_sum_m]
    for p_idx in range(num_participantes):
        col_l = get_column_letter(6 + p_idx)
        r_vals.append(f'=IF({col_l}{subtotal_row_m}="","",{col_l}{subtotal_row_m})')
    _write_row(ws, row, 1, r_vals,
               font=NORMAL_FONT, fill=FILL_LIGHT_BLUE,
               alignment=CENTER, number_format="0.0")
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center",
    )
    total_m_row = row
    row += 1

    # Fila: Total Juzgamiento
    r_vals_j = ["", "", "", "Total Juzgamiento (J)", pts_max_sum_j]
    for p_idx in range(num_participantes):
        col_l = get_column_letter(6 + p_idx)
        r_vals_j.append(f'=IF({col_l}{subtotal_row_j}="","",{col_l}{subtotal_row_j})')
    _write_row(ws, row, 1, r_vals_j,
               font=NORMAL_FONT, fill=FILL_LIGHT_ORANGE,
               alignment=CENTER, number_format="0.0")
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center",
    )
    total_j_row = row
    row += 1

    # Fila: Total Proceso
    r_vals_p = ["", "", "", "Total Proceso (P)", pts_max_sum_p]
    for p_idx in range(num_participantes):
        col_l = get_column_letter(6 + p_idx)
        r_vals_p.append(f'=IF({col_l}{subtotal_row_p}="","",{col_l}{subtotal_row_p})')
    _write_row(ws, row, 1, r_vals_p,
               font=NORMAL_FONT, fill=FILL_LIGHT_GREEN,
               alignment=CENTER, number_format="0.0")
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center",
    )
    total_p_row = row
    row += 1

    # Fila: PUNTAJE BRUTO TOTAL
    bruto_vals = ["", "", "", "PUNTAJE BRUTO TOTAL", pts_total_max]
    for p_idx in range(num_participantes):
        col_l = get_column_letter(6 + p_idx)
        formula = (
            f'=IF(COUNT({col_l}{first_data_row_m}:{col_l}{last_data_row_m},'
            f"{col_l}{first_data_row_j}:{col_l}{last_data_row_j},"
            f'{col_l}{first_data_row_p}:{col_l}{last_data_row_p})=0,"",'
            f"SUM({col_l}{total_m_row}:{col_l}{total_p_row}))"
        )
        bruto_vals.append(formula)
    for i, val in enumerate(bruto_vals):
        cell = ws.cell(row=row, column=1 + i, value=val)
        _apply_style(
            cell,
            font=Font(name="Calibri", bold=True, color=SENA_GREEN, size=11),
            fill=FILL_GOLD,
            alignment=CENTER,
            border=THIN_BORDER,
            number_format="0.0",
        )
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center",
    )
    bruto_row = row
    row += 1

    # Fila: PORCENTAJE FINAL (sobre 100)
    pct_vals = ["", "", "", "PORCENTAJE FINAL (sobre 100)", "100 %"]
    for p_idx in range(num_participantes):
        col_l = get_column_letter(6 + p_idx)
        formula = f'=IF({col_l}{bruto_row}="","",ROUND({col_l}{bruto_row}/{pts_total_max}*100,1))'
        pct_vals.append(formula)
    for i, val in enumerate(pct_vals):
        cell = ws.cell(row=row, column=1 + i, value=val)
        fmt = "0.0" if i >= 5 else None
        _apply_style(
            cell,
            font=Font(name="Calibri", bold=True, color=WHITE, size=12),
            fill=PatternFill("solid", fgColor=SENA_GREEN),
            alignment=CENTER,
            border=THIN_BORDER,
            number_format=fmt,
        )
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center",
    )
    pct_row = row
    row += 1

    # Fila: POSICIÓN
    pos_vals = ["", "", "", "POSICIÓN", ""]
    for p_idx in range(num_participantes):
        col_l = get_column_letter(6 + p_idx)
        formula = f'=IF({col_l}{pct_row}="","",RANK({col_l}{pct_row},$F${pct_row}:$H${pct_row},0))'
        pos_vals.append(formula)
    for i, val in enumerate(pos_vals):
        cell = ws.cell(row=row, column=1 + i, value=val)
        _apply_style(
            cell,
            font=Font(name="Calibri", bold=True, color=GOLD, size=14),
            fill=PatternFill("solid", fgColor=DARK_GREEN),
            alignment=CENTER,
            border=THIN_BORDER,
        )
    ws.cell(row=row, column=4).alignment = Alignment(
        horizontal="right", vertical="center",
    )
    row += 2

    # ===================================================================
    # OBSERVACIONES
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "OBSERVACIONES GENERALES",
        SECTION_FONT,
        FILL_HEADER_GREEN,
        CENTER,
    )
    row += 1
    _merge_and_style(
        ws, row, 1, row + 2, total_cols,
        "",
        NORMAL_FONT,
        FILL_WHITE,
        LEFT_TOP,
        THIN_BORDER,
    )
    row += 4

    # ===================================================================
    # FIRMAS DE JUECES
    # ===================================================================
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "FIRMAS DE LOS JUECES EVALUADORES",
        SECTION_FONT,
        FILL_HEADER_GREEN,
        CENTER,
    )
    row += 2

    # Distribuir las firmas horizontalmente
    firma_cols = [1, 4, 7]  # Columnas de inicio para cada firma
    firma_span = 2  # Cada firma ocupa 2 columnas

    for j_idx, juez in enumerate(JUECES):
        col_start = firma_cols[j_idx]
        col_end = col_start + firma_span - 1

        # Línea de firma
        _merge_and_style(
            ws, row, col_start, row, col_end,
            "_______________________________",
            Font(name="Calibri", size=10, color=MEDIUM_GRAY),
            FILL_WHITE,
            CENTER,
            Border(),
        )

    row += 1

    for j_idx, juez in enumerate(JUECES):
        col_start = firma_cols[j_idx]
        col_end = col_start + firma_span - 1

        # Nombre del juez
        _merge_and_style(
            ws, row, col_start, row, col_end,
            juez["nombre"],
            BOLD_FONT,
            FILL_WHITE,
            CENTER,
            Border(),
        )

    row += 1

    for j_idx, juez in enumerate(JUECES):
        col_start = firma_cols[j_idx]
        col_end = col_start + firma_span - 1

        # Cargo
        _merge_and_style(
            ws, row, col_start, row, col_end,
            juez["cargo"],
            SIGNATURE_FONT,
            FILL_WHITE,
            CENTER,
            Border(),
        )

    row += 2

    # Pie de página
    _merge_and_style(
        ws, row, 1, row, total_cols,
        "Plantilla de ejemplo para el reto de registro de pedidos — CGAO Skills 2026",
        Font(name="Calibri", size=8, italic=True, color="9E9E9E"),
        FILL_WHITE,
        CENTER,
        Border(),
    )

    # Guardar
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    project_root = Path(__file__).resolve().parent.parent
    dest = (
        project_root
        / "frontend"
        / "assets"
        / "documents"
        / "Evaluacion_Participantes_CGAO_Skills.xlsx"
    )
    result = generar_excel(dest)
    print(f"Plantilla Excel generada: {result}")
