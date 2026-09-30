import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from scripts.generar_excel_evaluacion import (
    _apply_style,
    _merge_and_style,
    _write_row,
    generar_excel,
)

ROOT_DIR = Path(__file__).resolve().parents[1]
PUBLISHED_WORKBOOK = ROOT_DIR / "frontend" / "assets" / "documents" / "Evaluacion_Participantes_CGAO_Skills.xlsx"
MASTER_RUBRIC = ROOT_DIR / "frontend" / "assets" / "documents" / "Rubrica_Maestra_Template.csv"


class RubricWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.workbook_path = Path(cls.temp_dir.name) / "evaluacion.xlsx"
        generar_excel(cls.workbook_path)
        cls.workbook = load_workbook(cls.workbook_path, data_only=False)
        cls.sheet = cls.workbook["Evaluación"]

    @classmethod
    def tearDownClass(cls):
        cls.workbook.close()
        cls.temp_dir.cleanup()

    def test_generated_criteria_match_the_master_rubric(self):
        ids = [
            self.sheet.cell(row=row, column=1).value
            for row in range(1, self.sheet.max_row + 1)
            if self.sheet.cell(row=row, column=1).value in {
                "M1", "M2", "M3", "M4", "M5", "M6", "M7",
                "J1", "J2", "J3", "P1", "P2",
            }
        ]
        self.assertEqual(ids, ["M1", "M2", "M3", "M4", "M5", "M6", "M7", "J1", "J2", "J3", "P1", "P2"])

        with MASTER_RUBRIC.open(encoding="utf-8-sig", newline="") as source:
            master_rows = list(csv.DictReader(source))
        generated_titles = {
            self.sheet.cell(row=row, column=1).value: self.sheet.cell(row=row, column=3).value
            for row in range(1, self.sheet.max_row + 1)
            if self.sheet.cell(row=row, column=1).value in ids
        }
        self.assertEqual(
            generated_titles,
            {criterion["ID"]: criterion["ASPECTO"] for criterion in master_rows},
        )

        max_points = {"M": 0, "J": 0, "P": 0}
        for row in range(1, self.sheet.max_row + 1):
            criterion_id = self.sheet.cell(row=row, column=1).value
            if criterion_id in ids:
                max_points[criterion_id[0]] += self.sheet.cell(row=row, column=5).value
        self.assertEqual(max_points, {"M": 69, "J": 23, "P": 8})

    def test_module_totals_and_final_formula_use_one_hundred_points(self):
        self.assertIn("Ejemplo formativo", self.sheet["A3"].value)
        self.assertIn("confirme condiciones con el instructor", self.sheet["A3"].value)
        labels = {}
        for row in range(1, self.sheet.max_row + 1):
            label = self.sheet.cell(row=row, column=4).value
            if label:
                labels[label] = row

        self.assertEqual(self.sheet.cell(labels["PUNTAJE BRUTO TOTAL"], 5).value, 100)
        self.assertEqual(self.sheet.cell(labels["Total Medición (M)"], 5).value, 69)
        self.assertEqual(self.sheet.cell(labels["Total Juzgamiento (J)"], 5).value, 23)
        self.assertEqual(self.sheet.cell(labels["Total Proceso (P)"], 5).value, 8)
        percentage_formula = self.sheet.cell(labels["PORCENTAJE FINAL (sobre 100)"], 6).value
        self.assertIn("/100", percentage_formula)

    def test_template_has_blank_inputs_and_weighted_judgment_formulas(self):
        criterion_rows = [
            row for row in range(1, self.sheet.max_row + 1)
            if self.sheet.cell(row=row, column=1).value in {"M1", "M2", "M3", "M4", "M5", "M6", "M7", "J1", "J2", "J3", "P1", "P2"}
        ]
        for row in criterion_rows:
            for column in (6, 7, 8):
                self.assertIsNone(self.sheet.cell(row=row, column=column).value)

        judgment_subtotal = next(
            row for row in range(1, self.sheet.max_row + 1)
            if self.sheet.cell(row=row, column=4).value == "SUBTOTAL JUZGAMIENTO (J)"
        )
        self.assertIn("SUMPRODUCT", self.sheet.cell(judgment_subtotal, 6).value)
        self.assertIn("/3", self.sheet.cell(judgment_subtotal, 6).value)

    def test_workbook_helpers_write_values_styles_and_merged_ranges(self):
        workbook = Workbook()
        sheet = workbook.active
        font = Font(name="Calibri", bold=True)
        fill = PatternFill("solid", fgColor="FFF8DC")
        alignment = Alignment(horizontal="center", vertical="center")

        _apply_style(sheet["A1"], font=font, fill=fill, alignment=alignment, number_format="0.0")
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00FFF8DC")
        self.assertEqual(sheet["A1"].number_format, "0.0")

        next_row = _write_row(sheet, 2, 2, ["M1", 7], fill=fill)
        self.assertEqual(next_row, 3)
        self.assertEqual([sheet.cell(row=2, column=column).value for column in (2, 3)], ["M1", 7])
        self.assertEqual(sheet["B2"].fill.fgColor.rgb, "00FFF8DC")

        _merge_and_style(sheet, 4, 1, 4, 2, "Título", font, fill, alignment)
        self.assertIn("A4:B4", {str(cell) for cell in sheet.merged_cells.ranges})
        self.assertEqual(sheet["A4"].value, "Título")
        self.assertTrue(sheet["A4"].font.bold)

    def test_published_workbook_matches_the_reproducible_rubric(self):
        workbook = load_workbook(PUBLISHED_WORKBOOK, data_only=False)
        try:
            sheet = workbook["Evaluación"]
            expected_ids = [
                "M1", "M2", "M3", "M4", "M5", "M6", "M7",
                "J1", "J2", "J3", "P1", "P2",
            ]
            criterion_rows = {
                sheet.cell(row=row, column=1).value: row
                for row in range(1, sheet.max_row + 1)
                if sheet.cell(row=row, column=1).value in expected_ids
            }
            self.assertEqual(list(criterion_rows), expected_ids)

            type_points = {"M": 0, "J": 0, "P": 0}
            for criterion_id, row in criterion_rows.items():
                type_points[criterion_id[0]] += sheet.cell(row=row, column=5).value
                for column in (6, 7, 8):
                    self.assertIn(sheet.cell(row=row, column=column).value, (None, ""))
            self.assertEqual(type_points, {"M": 69, "J": 23, "P": 8})

            criterion_text = {
                criterion_id: " ".join(
                    str(sheet.cell(row=criterion_rows[criterion_id], column=column).value or "")
                    for column in (3, 4)
                ).casefold()
                for criterion_id in ("M3", "M5", "M6", "M7", "P1", "P2")
            }
            self.assertRegex(criterion_text["M3"], r"alias|catálogo")
            self.assertRegex(criterion_text["M5"], r"almacenamiento local")
            self.assertRegex(criterion_text["M6"], r"orden descendente.*código menor")
            self.assertRegex(criterion_text["M7"], r"cierre y uso eficiente de la estación")
            self.assertRegex(criterion_text["M7"], r"archivos temporales")
            self.assertRegex(criterion_text["M7"], r"apaga el monitor o el equipo")
            self.assertRegex(criterion_text["P1"], r"ergonom|postura|orden")
            self.assertRegex(criterion_text["P2"], r"blanco para residuos aprovechables limpios y secos")
            self.assertRegex(criterion_text["P2"], r"verde para residuos orgánicos aprovechables")
            self.assertRegex(criterion_text["P2"], r"negro para residuos no aprovechables")

            module_d_points = sum(
                sheet.cell(row=row, column=5).value
                for row in criterion_rows.values()
                if sheet.cell(row=row, column=2).value.startswith("D ")
            )
            self.assertEqual(module_d_points, 10)

            validation_types = {
                validation.type
                for validation in sheet.data_validations.dataValidation
            }
            self.assertIn("custom", validation_types)
            self.assertIn("whole", validation_types)

            self.assertEqual(sheet.freeze_panes, "F1")
            self.assertEqual(sheet.sheet_view.zoomScale, 85)
            first_score_row = criterion_rows["M1"]
            self.assertEqual(sheet.cell(row=first_score_row, column=6).fill.fgColor.rgb[-6:], "FFF8DC")

            participant_header = next(
                row for row in range(1, sheet.max_row + 1)
                if sheet.cell(row=row, column=2).value == "Nombre completo"
            )
            for row in range(participant_header + 1, participant_header + 4):
                for column in (2, 3, 4):
                    self.assertIn(sheet.cell(row=row, column=column).value, (None, ""))

            all_values = " ".join(
                str(cell.value)
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
            self.assertNotRegex(all_values, r"María Paula|Juan Camilo|Andrés Felipe|2847391")
        finally:
            workbook.close()

    def test_published_copies_are_identical(self):
        copies = [
            PUBLISHED_WORKBOOK,
            ROOT_DIR / "frontend" / "assets" / "generated_docs" / PUBLISHED_WORKBOOK.name,
        ]
        source = PUBLISHED_WORKBOOK.read_bytes()
        for copy in copies[1:]:
            self.assertEqual(copy.read_bytes(), source, f"La copia {copy} debe estar sincronizada")


if __name__ == "__main__":
    unittest.main()
